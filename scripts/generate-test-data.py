#!/usr/bin/env python3
"""Generate reusable test fixtures; leave existing files untouched.

Use --tests-only to omit the multi-gigabyte development fixtures.
"""

import argparse
import os
import sqlite3
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
TESTDATA = ROOT / "testdata"
LARGE_FILES = Path(os.environ.get("OUTPUT_DIR", TESTDATA / "large-files"))


def ensure(path, write):
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    # Rename only after successful generation, so an interrupted run can retry.
    with tempfile.NamedTemporaryFile(prefix=f".{path.name}-", suffix=path.suffix, dir=path.parent, delete=False) as file:
        temporary = Path(file.name)
    try:
        write(temporary)
        if not path.exists():
            temporary.replace(path)
            print(f"Created {path}")
    finally:
        temporary.unlink(missing_ok=True)


def write_xlsx(path):
    workbook = Workbook(write_only=True)
    sales = workbook.create_sheet("Sales & Costs")
    sales.append(["name", "amount"])
    sales.append(["Alice", 10])
    sales.append(["Bob", 2])

    other = workbook.create_sheet("Other")
    other.append(["label"])
    other.append(["second"])
    for row in range(2, 151):
        other.append([f"row_{row}"])

    workbook.save(path)


def write_sqlite_tables(path):
    with sqlite3.connect(path) as connection:
        connection.executescript('CREATE TABLE metadata (name TEXT, amount INTEGER); CREATE TABLE "tile""data" (zoom INTEGER, tile_data BLOB);')
        connection.executemany("INSERT INTO metadata VALUES (?, ?)", [("Alice", 10), ("Bob", 2)])
        connection.executemany('INSERT INTO "tile""data" VALUES (?, ?)', [(1, bytes.fromhex("010203")), (2, b"a" * 1536), (3, b"a" * 1572864), (4, None)])


def write_paged_blobs(path):
    with sqlite3.connect(path) as connection:
        connection.execute("CREATE TABLE payload (id BIGINT, data BLOB)")
        blob = b"x" * 1048576
        connection.executemany("INSERT INTO payload VALUES (?, ?)", ((row, blob) for row in range(80)))


def write_text(path, text):
    path.write_text(text, encoding="utf-8")


def write_rows(path, count):
    with path.open("w", encoding="utf-8", newline="") as file:
        file.write("name,amount\n")
        for row in range(count):
            file.write(f"row_{row},{row}\n")


SORTED_PAGES_SQL = """CREATE TABLE data (name VARCHAR, amount INTEGER);
INSERT INTO data VALUES ('Zoe', 10), ('Alice', 10), ('Bob', 10), ('Carol', 2);
"""

FILTER_TABLE_SQL = r"""CREATE TABLE data ("a""b" VARCHAR, note VARCHAR, amount BIGINT, price DOUBLE, day DATE);
INSERT INTO data VALUES
    ('Alpha', 'first', 2, 1.5, DATE '2024-01-01'),
    ('a_b', 'second', 10, 2.5, DATE '2024-03-01'),
    ('Bob', 'alpha', 5, 3.5, DATE '2024-02-01'),
    ('b%', 'percent', 1, 4.5, DATE '2023-12-01'),
    ('x\back', 'slash', 9007199254740993, 5.5, DATE '2025-01-01'),
    (NULL, NULL, NULL, NULL, NULL);
"""

LITERAL_FILTERS_SQL = """CREATE TABLE data ("a""b" VARCHAR, amount INTEGER);
INSERT INTO data VALUES ('A_%', 10), ('aX%', 2), (NULL, NULL);
"""


def generate_test_fixtures():
    for filename, text in {
        "config.toml": '\ntheme = "light"\n',
        "inferred-types.csv": "name,amount,enabled\nAlice,10,true\nBob,2,false\n",
        "lazy.csv": "name,amount\nAlice,10\nBob,2\n",
        "lazy-appended.csv": "Carol,3\n",
        "sorted-pages.sql": SORTED_PAGES_SQL,
        "filter-table.sql": FILTER_TABLE_SQL,
        "literal-filters.sql": LITERAL_FILTERS_SQL,
    }.items():
        ensure(TESTDATA / filename, lambda path, text=text: write_text(path, text))
    ensure(TESTDATA / "multiple-sheets.xlsx", write_xlsx)
    ensure(TESTDATA / "sqlite-tables.db", write_sqlite_tables)
    ensure(TESTDATA / "page-cache.csv", lambda path: write_rows(path, 1024))
    ensure(TESTDATA / "eager.csv", lambda path: write_rows(path, 300))
    ensure(LARGE_FILES / "paged-blobs.db", write_paged_blobs)


def write_duckdb_sqlite(path, statements):
    # Keep the original DuckDB-based generator for the multi-gigabyte fixtures.
    sql_path = str(path).replace("'", "''")
    sql = f"ATTACH '{sql_path}' AS fixture (TYPE sqlite);\n{statements}\nDETACH fixture;"
    subprocess.run([os.environ.get("DUCKDB_BIN", "duckdb"), "-c", sql], check=True)


LARGE_DB_SQL = """
CREATE TABLE fixture.metadata (name TEXT, value TEXT);
INSERT INTO fixture.metadata VALUES
    ('purpose', 'large SQLite test fixture'),
    ('payload_rows', '1100'),
    ('payload_bytes', '1048576');
CREATE TABLE fixture.payload (id INTEGER PRIMARY KEY, data BLOB);
INSERT INTO fixture.payload
SELECT i, CAST(repeat(md5(CAST(i AS VARCHAR)), 32768) AS BLOB)
FROM range(1, 1101) t(i);
"""

WIDE_DB_SQL = """
CREATE TABLE fixture.records (
    id INTEGER PRIMARY KEY,
    customer_id BIGINT,
    active BOOLEAN,
    score DOUBLE,
    balance DECIMAL(12,2),
    signup_date DATE,
    last_seen TIMESTAMP,
    first_name VARCHAR,
    email VARCHAR,
    status VARCHAR,
    category VARCHAR,
    description VARCHAR,
    notes VARCHAR,
    region VARCHAR,
    priority SMALLINT,
    quantity INTEGER
);
INSERT INTO fixture.records
SELECT
    i,
    1000000000 + i,
    (i % 2 = 0),
    i * 0.125,
    CAST((i % 1000000) / 100.0 AS DECIMAL(12,2)),
    DATE '2020-01-01' + (i % 1826)::INTEGER,
    TIMESTAMP '2020-01-01 00:00:00' +
        (i % 31536000)::BIGINT * INTERVAL '1 second',
    'User ' || CAST(i AS VARCHAR),
    'user' || CAST(i AS VARCHAR) || '@example.test',
    CASE i % 4
        WHEN 0 THEN 'new'
        WHEN 1 THEN 'active'
        WHEN 2 THEN 'paused'
        ELSE 'closed'
    END,
    CASE i % 3
        WHEN 0 THEN 'alpha'
        WHEN 1 THEN 'beta'
        ELSE 'gamma'
    END,
    CAST(repeat('Description text for test row ' || CAST(i AS VARCHAR) || '. ', 64) AS VARCHAR),
    CAST(repeat('Additional notes for wide SQLite fixture row ' || CAST(i AS VARCHAR) || '. ', 48) AS VARCHAR),
    CASE i % 5
        WHEN 0 THEN 'north'
        WHEN 1 THEN 'south'
        WHEN 2 THEN 'east'
        WHEN 3 THEN 'west'
        ELSE 'central'
    END,
    CAST(i % 10 AS SMALLINT),
    CAST(i % 1000 AS INTEGER)
FROM range(1, 300001) t(i);
"""


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--tests-only", action="store_true", help="omit multi-gigabyte development fixtures")
    args = parser.parse_args()
    generate_test_fixtures()
    if not args.tests_only:
        ensure(LARGE_FILES / "large-test.db", lambda path: write_duckdb_sqlite(path, LARGE_DB_SQL))
        ensure(LARGE_FILES / "large-test-columns.db", lambda path: write_duckdb_sqlite(path, WIDE_DB_SQL))


if __name__ == "__main__":
    main()
